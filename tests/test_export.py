"""Excel экспорт — /api/member/export, /api/organization/export, /api/admin/forms/<id>/results/export."""
import datetime as dt
import io
import os

import openpyxl

import core.xlsx as xlsx
from conftest import uniq
from _forms_helpers import published_form, submit, opt

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _book(r):
    assert r.status_code == 200, r.get_data(as_text=True)[:300]
    assert r.mimetype == MIME
    return openpyxl.load_workbook(io.BytesIO(r.data))


def _org(api, **kw):
    body = {"name": uniq("Экспорт сургууль"), "org_code": f"{int(uniq('')) % 900 + 100:03d}",
            "school_category_id": 12}
    body.update(kw)
    r = api.post("/api/organization", json=body)
    assert r.status_code == 201, r.get_json()
    return r.get_json()


def _member(api, org_id, **kw):
    body = {"organization_id": org_id, "last_name": "Бат", "first_name": uniq("Дорж")}
    body.update(kw)
    r = api.post("/api/member", json=body)
    assert r.status_code == 201, r.get_json()
    return r.get_json()


# ================================ member ================================
def test_member_export_layout_and_values(api):
    org = _org(api)
    m = _member(api, org["id"], birth_date="1990-05-06", gender="эм", register_number="УБ90050612",
                union_joined_date="2015-03-01", member_status="Идэвхтэй", status="баталгаажсан")
    _member(api, org["id"])                                   # хоосон талбартай гишүүн
    api.post("/api/contact", json={"owner_type": "member", "owner_id": m["id"], "type": "утас",
                                   "value": "99112233"})
    api.post("/api/contact", json={"owner_type": "member", "owner_id": m["id"], "type": "факс",
                                   "value": "70112233"})
    r = api.get(f"/api/member/export?organization_id={org['id']}")
    assert r.headers["Content-Disposition"].startswith("attachment")
    assert f"{dt.date.today().isoformat()}.xlsx" in r.headers["Content-Disposition"]
    assert "filename*=UTF-8''" in r.headers["Content-Disposition"]          # кирилл нэр
    ws = _book(r).active
    header = [c.value for c in ws[1]]
    assert header == ["№", "Овог нэр", "Төрсөн он", "Хүйс", "Регистрийн дугаар",
                      "ҮЭ-ийн бүртгэлийн дугаар", "ҮЭ-д элссэн огноо", "ҮЭ-ийн гишүүний статус",
                      "Статус", "Албан тушаал", "Мэргэжил", "Утасны дугаар", "Байгууллага"]
    assert all(c.font.b for c in ws[1]) and ws.freeze_panes == "A2"
    assert ws.max_row == 3                                    # толгой + 2 гишүүн
    row = [c.value for c in ws[2]]
    assert row[0] == 1 and row[1] == f"Бат {m['first_name']}" and row[2] == 1990
    assert row[3] == "эм" and row[4] == "УБ90050612"
    assert row[6] == dt.datetime(2015, 3, 1) and ws.cell(2, 7).is_date      # огноо нүд
    assert row[7:9] == ["Идэвхтэй", "баталгаажсан"]
    assert row[11] == "99112233, 70112233" and row[12] == org["name"]
    empty = [c.value for c in ws[3]]
    assert empty[2] is None and empty[6] is None and "null" not in [str(v) for v in empty]


def test_member_export_matches_list_filters(api):
    org = _org(api)
    for active in (1, 1, 0):
        _member(api, org["id"], is_active=active)
    listed = api.get(f"/api/member?organization_id={org['id']}&is_active=1").get_json()
    ws = _book(api.get(f"/api/member/export?organization_id={org['id']}&is_active=1")).active
    assert ws.max_row - 1 == len(listed) == 2


def test_member_export_respects_scope(api, make_user):
    inside, outside = _org(api), _org(api)
    _member(api, inside["id"]), _member(api, outside["id"])
    mgr, user = make_user(["member.read", "organization.read"])
    assert api.put(f"/api/user/{user['id']}/scope",
                   json={"organization_id": inside["id"]}).status_code == 200
    ws = _book(mgr.get("/api/member/export")).active
    orgs = {ws.cell(r, 13).value for r in range(2, ws.max_row + 1)}
    assert orgs == {inside["name"]}
    ws = _book(mgr.get("/api/organization/export")).active
    assert [ws.cell(r, 3).value for r in range(2, ws.max_row + 1)] == [inside["name"]]


# ============================= organization =============================
def test_org_export_layout_and_stats(api):
    org = _org(api, registration_number="1234567", contact_name="Захирал", phone1="99001122",
               email="a@b.mn")
    _member(api, org["id"], gender="эм", birth_date="2000-01-01")
    _member(api, org["id"], gender="эр", birth_date="1970-01-01")
    ws = _book(api.get("/api/organization/export?school_category_id=12")).active
    assert [c.value for c in ws[1]] == ["№", "Код", "Нэр", "Төрөл", "Регистрийн дугаар",
                                        "Удирдлагын мэдээлэл", "Утас1", "Утас2", "И-мэйл",
                                        "Гишүүдийн тоо", "Эмэгтэй гишүүд", "35 хүртэлх насны гишүүд"]
    assert ws.freeze_panes == "A2"
    row = next([c.value for c in r] for r in ws.iter_rows(min_row=2) if r[2].value == org["name"])
    assert row[1] == org["full_code"] and row[3] == org["school_category_short_name"]
    assert row[4:9] == ["1234567", "Захирал", "99001122", None, "a@b.mn"]
    assert row[9:12] == [2, 1, 1]
    listed = api.get("/api/organization?school_category_id=12").get_json()
    assert ws.max_row - 1 == len(listed)


# ============================ survey results ============================
def test_form_results_export(api, anon):
    fid, qs = published_form(api, title=uniq("Экспорт судалгаа"))
    r1 = submit(anon, fid, [{"question_id": qs["single"]["id"], "option_ids": [opt(qs["single"], 0)]},
                            {"question_id": qs["multi"]["id"],
                             "option_ids": [opt(qs["multi"], 0), opt(qs["multi"], 2)]},
                            {"question_id": qs["scale"]["id"], "numeric_value": 4},
                            {"question_id": qs["text"]["id"], "text_value": "Сайн"}])
    r2 = submit(anon, fid, [{"question_id": qs["single"]["id"], "option_ids": [opt(qs["single"], 1)]}])
    assert r1.status_code == 201 and r2.status_code == 201, (r1.get_json(), r2.get_json())
    r = api.get(f"/api/admin/forms/{fid}/results/export")
    assert "%D2%AF%D1%80-%D0%B4%D2%AF%D0%BD" in r.headers["Content-Disposition"]   # "үр-дүн"
    wb = _book(r)
    assert wb.sheetnames == ["Хариултууд", "Дүгнэлт"]
    ws = wb["Хариултууд"]
    header = [c.value for c in ws[1]]
    assert header[:2] == ["№", "Огноо"] and len(header) == 6
    assert "Хэрэглэгч" not in header                           # хариулагч тодорхойлогдохгүй
    first = [c.value for c in ws[2]]
    assert first[0] == 1 and isinstance(first[1], dt.datetime)
    assert first[2] == "А" and first[3] == "А, В" and first[4] == 4 and first[5] == "Сайн"
    second = [c.value for c in ws[3]]
    assert second[2] == "Б" and second[3:] == [None, None, None]
    summary = [[c.value for c in r] for r in wb["Дүгнэлт"].iter_rows(min_row=2)]
    single = [s for s in summary if s[0] == qs["single"]["title"]]
    assert [(s[1], s[2], s[3]) for s in single] == [("А", 1, 50), ("Б", 1, 50), ("В", 0, 0)]
    scale = [s for s in summary if s[0] == qs["scale"]["title"]]
    assert scale[0][1:3] == ["Дундаж", 1] and scale[0][4] == 4
    assert [(s[1], s[2]) for s in scale[1:]] == [("1", 0), ("2", 0), ("3", 0), ("4", 1), ("5", 0)]


def test_empty_form_export_has_headers_only(api):
    fid, _ = published_form(api)
    wb = _book(api.get(f"/api/admin/forms/{fid}/results/export"))
    assert wb["Хариултууд"].max_row == 1


# ================================ errors ================================
def test_errors_are_json(api, anon, make_user):
    r = api.get("/api/admin/forms/999999/results/export")
    assert r.status_code == 404 and "error" in r.get_json()
    assert anon.get("/api/member/export").status_code == 401
    u, _ = make_user(["member.read"])
    r = u.get("/api/organization/export")
    assert r.status_code == 403 and "error" in r.get_json()
    assert api.get("/api/member/export?page=1").status_code == 200          # хуудаслалтгүй


def test_temp_file_removed_after_response(api, monkeypatch):
    made = []
    real = xlsx.tempfile.mkstemp
    monkeypatch.setattr(xlsx.tempfile, "mkstemp",
                        lambda **kw: made.append(real(**kw)) or made[-1])
    r = api.get("/api/member/export")
    assert r.status_code == 200 and r.data[:2] == b"PK"      # xlsx = zip
    assert made and not os.path.exists(made[0][1])          # хариу буцмагц нэр нь алга
